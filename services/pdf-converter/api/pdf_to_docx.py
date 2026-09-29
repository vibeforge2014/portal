from __future__ import annotations

import sys

import fitz
from docx import Document
from pdf2docx import Converter


def compact(value: str) -> str:
    return "".join(value.lower().split())


def append_ocr_text(source: str, destination: str, page_indexes: list[int]) -> None:
    if not page_indexes:
        return
    document = Document(destination)
    existing = compact("\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    ))
    source_pdf = fitz.open(source)
    added = False
    try:
        for page_index in page_indexes:
            lines = [line.strip() for line in source_pdf[page_index].get_text("text", sort=True).splitlines() if line.strip()]
            recognized = compact(" ".join(lines))
            if not recognized or (len(recognized) >= 12 and recognized[:24] in existing):
                continue
            document.add_page_break()
            for line in lines:
                document.add_paragraph(line)
            existing += recognized
            added = True
    finally:
        source_pdf.close()
    if added:
        document.save(destination)


def main() -> None:
    source, destination = sys.argv[1:3]
    page_indexes = [int(value) for value in sys.argv[3].split(",") if value] if len(sys.argv) > 3 else []
    converter = Converter(source)
    try:
        converter.convert(destination, multi_processing=False)
    finally:
        converter.close()
    append_ocr_text(source, destination, page_indexes)


if __name__ == "__main__":
    main()
